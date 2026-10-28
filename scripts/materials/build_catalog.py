"""
Builds src/lib/materials/catalog.ts and the AI's reference list from the principal's material list
(委托方提供的《中国日用商品内外包装材料总表》), kept as it was received in data/materials/:

    python3 scripts/materials/build_catalog.py [path to the .xlsx]

Three sheets are read: 材料总表 (168 materials), 合规路径 (23 rules) and 设计使用说明 (usage notes).
Each table starts at the row whose first cell is 编号 or 事项. The Chinese is copied as it is;
English comes from en_materials.py and en_rules.py, which must cover every row (the script stops
if one is missing, so a new list can't slip through half translated). Needs openpyxl.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from en_materials import DOCS, EN, FAMILIES, PARTS, SOURCES  # noqa: E402
from en_rules import NOTES, PATHS, RULES  # noqa: E402

SRC = ROOT / "data" / "materials" / "china-packaging-materials-2026-10-28.xlsx"
OUT = ROOT / "src" / "lib" / "materials" / "catalog.ts"
IDS = ROOT / "src" / "lib" / "materials" / "ids.ts"
REF = ROOT / "skills" / "packaging-visual-direction" / "references" / "material-list.md"

FAMILY_IDS = {
    "纸与纸板": "paper",
    "纸基面料与功能纸": "paperSpecial",
    "塑料与塑料形态": "plastic",
    "复合与阻隔结构": "laminate",
    "金属": "metal",
    "玻璃与陶瓷": "glassCeramic",
    "木竹与天然结构材料": "wood",
    "织物与皮革": "textile",
    "缓冲与定位材料": "cushion",
    "密封与包装辅料": "consumable",
    "其他有既有应用的材料": "other",
}

# rows the list itself limits to non-food uses (its boundary or typical use says so)
NOT_FOOD = re.compile(r"非食品|不直接接触食品|不接触食品|无食品接触|不得直接作为食品|不用于食品|不能直接当食品|不作食品")
# the boundary says how food contact is reviewed, so the typical use alone doesn't rule it out
FOOD_PATH_GIVEN = {"M164"}
# the boundary rules out cosmetics contact as well
NOT_COSMETICS = {"M069"}


def table(ws, first):
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r and r[0] == first)
    head = [str(h).strip() if h is not None else "" for h in rows[start]]
    out = []
    for r in rows[start + 1:]:
        if not r or r[0] is None or str(r[0]).strip() == "":
            continue
        out.append({h: ("" if v is None else str(v).strip()) for h, v in zip(head, r) if h})
    return out


def codes(text, prefix="F"):
    return sorted(set(re.findall(prefix + r"\d\d", text)))


def parse_path(path, boundary, name):
    """The rules a material brings, by the condition that triggers them."""
    food, courier, bag, g02 = set(), False, False, False
    for tok in [t.strip() for t in path.split("；") if t.strip()]:
        # "F04或F06": glass takes the glass rule, ceramics the ceramic one
        alt = re.search(r"(F\d\d)或(F\d\d)", tok)
        if alt:
            a, b = alt.groups()
            chosen = [c for c in (a, b) if c in boundary] or [a if "玻璃" in name else b]
            tok = tok.replace(alt.group(0), "/".join(chosen))
        if tok == "G02":
            g02 = True
        elif tok.startswith("G03"):
            courier = True
        elif tok.startswith("P01"):
            bag = True
        elif "食品接触" in tok or "涂层" in tok or "按用途" in tok:
            food |= set(codes(tok)) | {"F01"}
    # the boundary names the rules for this row's own material (硅橡胶 F12, 搪瓷 F14, …)
    food |= set(codes(boundary))
    if re.search(r"G03", boundary):
        courier = True
    if re.search(r"P01", boundary):
        bag = True
    special = codes(boundary, "D") + codes(boundary, "M")
    return sorted(food), courier, bag, g02, special


def ts(v):
    return json.dumps(v, ensure_ascii=False)


def pair(zh, en):
    return {"zh": zh, "en": en}


def main(src):
    wb = openpyxl.load_workbook(src)
    mats = table(wb["材料总表"], "编号")
    rules = table(wb["合规路径"], "编号")
    notes = table(wb["设计使用说明"], "事项")
    title = wb["材料总表"].cell(row=2, column=1).value
    assert len(mats) == 168, len(mats)

    missing = [m["编号"] for m in mats if m["编号"] not in EN]
    missing += [r["编号"] for r in rules if r["编号"] not in RULES]
    missing += [n["事项"] for n in notes if n["事项"] not in NOTES]
    missing += [m["合规路径编号"] for m in mats if m["合规路径编号"] not in PATHS]
    if missing:
        sys.exit(f"no English for: {missing}")

    docs = sorted(set(m["采购核对资料"] for m in mats))
    sources = sorted(set(m["品类资料来源"] for m in mats))
    paths = sorted(set(m["合规路径编号"] for m in mats))

    out = []
    for m in mats:
        i = m["编号"]
        name, use, notes_en, boundary_en = EN[i]
        food, courier, bag, g02, special = parse_path(m["合规路径编号"], m["使用边界"], m["材料或品种"])
        no_food = i not in FOOD_PATH_GIVEN and bool(NOT_FOOD.search(m["使用边界"] + " " + m["典型用途"]))
        url = m.get("资料链接", "")
        url = url if url.startswith("http") else ""
        rec = {
            "id": i,
            "family": FAMILY_IDS[m["材料大类"]],
            "name": pair(m["材料或品种"], name),
            "trade": m["采购常用名称"],
            "parts": pair(m["包装部位"], PARTS[m["包装部位"]]),
            "use": pair(m["典型用途"], use),
            "notes": pair(m["设计与加工要点"], notes_en),
            "boundary": pair(m["使用边界"], boundary_en),
            "path": paths.index(m["合规路径编号"]),
            "food": food,
            "docs": docs.index(m["采购核对资料"]),
            "source": sources.index(m["品类资料来源"]),
        }
        if courier:
            rec["courier"] = True
        if bag:
            rec["bag"] = True
        if g02:
            rec["g02"] = True
        if special:
            rec["special"] = special
        if url:
            rec["url"] = url
        if no_food:
            rec["noFood"] = True
        if i in NOT_COSMETICS:
            rec["noCosmetics"] = True
        out.append(rec)

    rule_out = []
    for r in rules:
        scope, std, status, needs, rdocs, checked = RULES[r["编号"]]
        rule_out.append({
            "id": r["编号"],
            "scope": pair(r["适用范围"], scope),
            "standard": pair(r["标准或法规"], std),
            "status": pair(r["状态/核验说明"], status),
            "needs": pair(r["需要满足的条件"], needs),
            "docs": pair(r["应取得或确认的资料"], rdocs),
            "checked": pair(r["来源核验层级"], checked),
            "urls": [u.strip() for u in r["来源链接"].split("\n") if u.strip().startswith("http")],
        })
    note_out = [{"topic": pair(n["事项"], NOTES[n["事项"]][0]), "text": pair(n["说明"], NOTES[n["事项"]][1])} for n in notes]

    lines = [
        "/**",
        " * The principal's material list (委托方提供的材料总表), generated by scripts/materials/build_catalog.py",
        " * from data/materials/china-packaging-materials-2026-10-28.xlsx. Do not edit by hand: change the",
        " * list or the translations and run the script again.",
        " */",
        'import type { Pair } from "../i18n";',
        "",
        f"export const LIST = {ts({'title': pair(title, 'Packaging materials for everyday consumer goods in China'), 'checked': '2026-10-28', 'file': 'data/materials/china-packaging-materials-2026-10-28.xlsx', 'count': len(out)})} as const;",
        "",
        "export type FamilyId = " + " | ".join(f'"{v}"' for v in FAMILY_IDS.values()) + ";",
        "export const FAMILIES: { id: FamilyId; label: Pair }[] = " + ts([{"id": FAMILY_IDS[k], "label": pair(k, FAMILIES[k])} for k in FAMILY_IDS]) + ";",
        "",
        "export type Material = {",
        "  id: string;",
        "  family: FamilyId;",
        "  name: Pair;",
        "  /** 采购常用名称: what suppliers call it (Chinese, as listed) */",
        "  trade: string;",
        "  parts: Pair;",
        "  use: Pair;",
        "  /** 设计与加工要点 */",
        "  notes: Pair;",
        "  /** 使用边界: the limits of use the list sets */",
        "  boundary: Pair;",
        "  /** index into PATHS: the compliance path as written */",
        "  path: number;",
        "  /** the food-contact rules (F…) for direct food contact */",
        "  food: string[];",
        "  /** the path adds G03 when the material is used in courier packaging */",
        "  courier?: true;",
        "  /** the path adds P01 for shopping bags or other designated plastic items */",
        "  bag?: true;",
        "  /** the path lists G02 (food and cosmetics sales packaging) */",
        "  g02?: true;",
        "  /** special routes the boundary points to (D01 medicines, M01 medical devices) */",
        "  special?: string[];",
        "  /** index into DOCS: the supplier documents to check when buying */",
        "  docs: number;",
        "  /** index into SOURCES */",
        "  source: number;",
        "  url?: string;",
        "  /** the list limits this row to uses that don't touch food */",
        "  noFood?: true;",
        "  /** …or cosmetics */",
        "  noCosmetics?: true;",
        "};",
        "",
        "export type Rule = { id: string; scope: Pair; standard: Pair; status: Pair; needs: Pair; docs: Pair; checked: Pair; urls: string[] };",
        "",
        "export const PATHS: Pair[] = " + ts([pair(p, PATHS[p]) for p in paths]) + ";",
        "export const DOCS: Pair[] = " + ts([pair(d, DOCS[d]) for d in docs]) + ";",
        "export const SOURCES: Pair[] = " + ts([pair(s, SOURCES[s]) for s in sources]) + ";",
        "",
        "export const MATERIALS: Material[] = [",
        *[f"  {ts(r)}," for r in out],
        "];",
        "",
        "export const RULES: Rule[] = [",
        *[f"  {ts(r)}," for r in rule_out],
        "];",
        "",
        "export const NOTES: { topic: Pair; text: Pair }[] = " + ts(note_out) + ";",
        "",
    ]
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(lines), "utf-8")
    # the numbers alone, for checks that should not load the whole list (the board validator)
    IDS.write_text(
        "/** The material numbers on the principal's list (generated by scripts/materials/build_catalog.py). */\n"
        f"export const MATERIAL_IDS: ReadonlySet<string> = new Set({ts([r['id'] for r in out])});\n",
        "utf-8",
    )

    # the AI's copy: every name it may use, with the limits that decide where it can go
    md = [
        "---",
        "name: material-list",
        "description: The principal's material list. Name packaging materials only from this list, by number.",
        "---",
        "",
        f"# 材料清单（委托方提供，核对日期 2026-10-28，共 {len(out)} 项）",
        "",
        "提到包装材料时只能使用下表中的材料，写明编号，例如“单面铜版纸（M019）”。表外材料一律不写；需要表外材料时，写“表内暂无合适材料，请与委托方确认”。",
        "每项只在“使用边界”内使用。直接接触食品或化妆品的部件必须选未标“非食品”的材料。",
        "",
        "| 编号 | 材料 | 大类 | 典型用途 | 使用边界 |",
        "|---|---|---|---|---|",
        *[f"| {m['编号']} | {m['材料或品种']} | {m['材料大类']} | {m['典型用途']} | {m['使用边界']} |" for m in mats],
        "",
    ]
    REF.write_text("\n".join(md), "utf-8")
    print(f"{len(out)} materials, {len(rule_out)} rules, {len(note_out)} notes -> {OUT.relative_to(ROOT)}, {REF.relative_to(ROOT)}")
    print("not for food:", sum(1 for r in out if r.get("noFood")), "with links:", sum(1 for r in out if r.get("url")))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else SRC)
